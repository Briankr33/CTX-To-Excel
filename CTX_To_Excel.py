import re
import glob
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

#Create Excel workbook and assign variable for active sheet
wb = Workbook()
ws1 = wb.active
ws1.append(['I/O', 'Device ID', 'Description', 'File Name'])

#Define regex patterns
regexGroup = r"GmmiLinkContextFrameObject\s+.+\s+.+\s+.+\s+.+\s+.+\s+.+\s+.+\s+.+\s+.+\s+.+"
regexBit = r"%?[I|Q]\d{3,4}"
regexDevice = r"GmmiObject\s\"\D+\d{5}\""
regexDescription = r"\"Desc2\"\s\".+\""

#Find and open all .ctx files in current directory
file_list = glob.glob("*.ctx")
for file in file_list:
    print(file)

    #Open .ctx file and read contents
    with open(file, 'r',encoding= 'utf_16') as f:
        readList = f.read()

    #Declare and initialize variable for list of regex matches
    matchList = re.findall(regexGroup, readList)

    listLength = len(matchList)

    #Loop through the list of matches and find device info
    for _ in range(listLength):
        rowNumber = ws1.max_row + 1
        columnLetter = 'A'
        matchItem1 = re.findall(regexBit, matchList[_])
        if matchItem1:
            #Some devices contain multiple bit addresses
            if len(matchItem1) > 1:
                valueColumnA = ", ".join(matchItem1)
            else:
                valueColumnA = matchItem1[0]
            print(valueColumnA)
            ws1["%s%s" % (columnLetter, rowNumber)] = valueColumnA

        columnLetter = 'B'
        matchItem2 = re.findall(regexDevice, matchList[_])
        if matchItem2:
            valueColumnB = matchItem2[0]
            ws1["%s%s" % (columnLetter, rowNumber)] = valueColumnB

        columnLetter = 'C'
        matchItem3 = re.findall(regexDescription, matchList[_])
        if matchItem3:
            valueColumnC = matchItem3[0]
            ws1["%s%s" % (columnLetter, rowNumber)] = valueColumnC

        columnLetter = 'D'
        if matchItem1 or matchItem2 or matchItem3:
            ws1["%s%s" % (columnLetter, rowNumber)] = file

tab = Table(displayName= 'Table1', ref='A1:%s%s' % (columnLetter, rowNumber))
style = TableStyleInfo(name='TableStyleLight1', showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
tab.tableStyleInfo = style
ws1.add_table(tab)

#Save excel workbook
print("Code finished successfully.")
wb.save('Practice.xlsx')
print("Workbook saved.")
